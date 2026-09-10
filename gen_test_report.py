"""
生成中国象棋全量测试报告 (.docx)
"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_NAME = 'Microsoft YaHei'

def _set_rFonts(element, font_name):
    rPr = element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        key = qn(attr)
        if rFonts.get(key) is not None:
            del rFonts.attrib[key]

def set_doc_default_font(doc, font_name=FONT_NAME):
    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn('w:docDefaults'))
    if doc_defaults is None:
        doc_defaults = OxmlElement('w:docDefaults')
        styles_el.insert(0, doc_defaults)
    rPr_default = doc_defaults.find(qn('w:rPrDefault'))
    if rPr_default is None:
        rPr_default = OxmlElement('w:rPrDefault')
        doc_defaults.append(rPr_default)
    rPr = rPr_default.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        rPr_default.append(rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        key = qn(attr)
        if rFonts.get(key) is not None:
            del rFonts.attrib[key]

def apply_chinese_fonts(doc, font_name=FONT_NAME):
    set_doc_default_font(doc, font_name)
    _set_rFonts(doc.styles['Normal'].element, font_name)
    for i in range(1, 10):
        try:
            _set_rFonts(doc.styles[f'Heading {i}'].element, font_name)
        except KeyError:
            pass

def set_run_font(run, font_name=FONT_NAME, size=None, bold=False, color=None):
    run.font.name = font_name
    _set_rFonts(run._element, font_name)
    if size: run.font.size = size
    if bold: run.font.bold = True
    if color: run.font.color.rgb = color

def add_para(doc, text, size=Pt(12), bold=False, color=None, alignment=None, space_after=Pt(6)):
    p = doc.add_paragraph()
    if alignment: p.alignment = alignment
    p.paragraph_format.space_after = space_after
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p

def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=Pt(16 if level==1 else 14 if level==2 else 12), bold=True)
    return h

def set_cell(cell, text, size=Pt(9), bold=False, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align: p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold)

def shade_cell(cell, color):
    s = OxmlElement('w:shd')
    s.set(qn('w:fill'), color)
    s.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(s)

def make_table(doc, headers, rows):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, size=Pt(9), bold=True)
        shade_cell(table.rows[0].cells[i], 'D5E8F0')
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell(table.rows[r+1].cells[c], val, size=Pt(9))
    return table

doc = Document()
apply_chinese_fonts(doc)
section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(3.17)
section.right_margin = Cm(3.17)

# ========== 封面 ==========
for _ in range(6):
    doc.add_paragraph()
add_para(doc, '中国象棋（多模式）', size=Pt(28), bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, '', size=Pt(12))
add_para(doc, '全量测试报告', size=Pt(22), bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x2E, 0x75, 0xB6))
add_para(doc, '', size=Pt(12))
add_para(doc, '测试日期：2026-09-04', size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, '测试人：Luis（露易丝）', size=Pt(14), alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, '', size=Pt(12))
add_para(doc, '项目路径：workCodes/中国象棋/index.html', size=Pt(12), alignment=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ========== 一、测试概述 ==========
add_heading_styled(doc, '一、测试概述', level=1)

add_heading_styled(doc, '1.1 测试对象', level=2)
make_table(doc,
    ['项目', '说明'],
    [
        ['项目名称', '中国象棋（多模式对战）'],
        ['项目路径', 'workCodes/中国象棋/index.html'],
        ['技术形态', '单文件纯前端（HTML + Canvas + 原生 JS）'],
        ['包含模式', '普通局 / ban局 / ban进阶 / 心腹局 / 内臣局'],
        ['AI 能力', '5 档难度（新手~宗师）+ 思考时间可调'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '1.2 测试环境', level=2)
make_table(doc,
    ['项目', '说明'],
    [
        ['本地服务', 'python -m http.server 8765'],
        ['页面地址', 'http://localhost:8765/index.html'],
        ['浏览器', 'Google Chrome（headless=new 模式）'],
        ['驱动方式', 'Chrome DevTools Protocol（CDP）'],
        ['逻辑执行', 'Node.js v24.14.0 + vm 沙箱'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '1.3 测试方法', level=2)
add_para(doc, '采用三层测试策略：')
make_table(doc,
    ['层次', '方法', '覆盖范围'],
    [
        ['逻辑层', '提取核心 JS 逻辑（ChessBoard/ChessAI/MODES）用 Node vm 沙箱测试', '棋盘规则、走法、吃子、将杀、AI 搜索、模式核心逻辑'],
        ['交互层', '提取控制器逻辑（onBoardClick/悔棋/AI 回调）模拟测试', '竞态控制（epoch）、悔棋边界、对局流程、AI 禁子/选心腹'],
        ['E2E 层', 'Chrome headless + CDP 真实驱动页面', '首屏渲染、真实点击走棋、模式切换、弹窗、悔棋、重开、竞态、控制台错误'],
    ]
)
add_para(doc, '', size=Pt(6))

doc.add_page_break()

# ========== 二、测试结果总览 ==========
add_heading_styled(doc, '二、测试结果总览', level=1)

make_table(doc,
    ['测试层次', '通过', '失败', '总计', '通过率'],
    [
        ['逻辑层（test_logic.js）', '29', '0', '29', '100%'],
        ['交互层（test_controller.js）', '14', '0', '14', '100%'],
        ['E2E 层（test_e2e.js）', '14', '0', '14', '100%'],
        ['E2E 补充（test_banpro_e2e.js）', '3', '0', '3', '100%'],
        ['合计', '60', '0', '60', '100%'],
    ]
)

add_para(doc, '', size=Pt(6))
add_para(doc, '测试结论：60/60 全部通过，通过率 100%。', size=Pt(12), bold=True, color=RGBColor(0x00, 0x80, 0x00))

doc.add_page_break()

# ========== 三、逻辑层测试 ==========
add_heading_styled(doc, '三、逻辑层测试详情（29 用例）', level=1)

add_heading_styled(doc, '3.1 棋盘初始化（5 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['L-01', '初始棋子数量 = 32', '✅ 通过'],
        ['L-02', '红方先行', '✅ 通过'],
        ['L-03', '双方将帅位置正确', '✅ 通过'],
        ['L-04', '兵卒位置正确', '✅ 通过'],
        ['L-05', '炮位置正确', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '3.2 基本走棋规则（7 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['L-06', '兵初始只能向前走 1 格', '✅ 通过'],
        ['L-07', '兵过河后可以左右走', '✅ 通过'],
        ['L-08', '马走日且不被蹩马腿', '✅ 通过'],
        ['L-09', '车直线走棋', '✅ 通过'],
        ['L-10', '炮隔山打子', '✅ 通过'],
        ['L-11', '将帅不能走出九宫格', '✅ 通过'],
        ['L-12', '不能走出送将的棋', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '3.3 吃子与将杀（4 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['L-13', '吃子后 captured 返回正确', '✅ 通过'],
        ['L-14', 'isCheckmate 正常对局返回 false', '✅ 通过'],
        ['L-15', 'isCheckmate 无子可动返回 true', '✅ 通过'],
        ['L-16', 'AI 搜索不崩溃并返回合法走法（BUG-2 核心验证）', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '3.4 模式逻辑（9 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['L-17', '经典模式 init 设置 phase=play', '✅ 通过'],
        ['L-18', 'ban 局 init 进入选禁阶段', '✅ 通过'],
        ['L-19', 'ban 局点击非红方棋子被忽略', '✅ 通过'],
        ['L-20', 'ban 局红方禁子后进入 AI 禁', '✅ 通过'],
        ['L-21', '心腹局 init 进入选心腹阶段', '✅ 通过'],
        ['L-22', '心腹局选择红方棋子', '✅ 通过'],
        ['L-23', '心腹局将死时心腹牺牲', '✅ 通过'],
        ['L-24', '内臣局 init 进入选内臣阶段', '✅ 通过'],
        ['L-25', '内臣局吃将后内臣继位（BUG-3 核心验证）', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '3.5 悔棋与 AI（4 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['L-26', 'undo 回退一步', '✅ 通过'],
        ['L-27', '空历史 undo 返回 false', '✅ 通过'],
        ['L-28', 'AI 各难度搜索不崩溃（深度 1~5）', '✅ 通过'],
        ['L-29', 'AI 在吃子场景优先吃子', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

doc.add_page_break()

# ========== 四、交互层测试 ==========
add_heading_styled(doc, '四、交互层测试详情（14 用例）', level=1)

add_heading_styled(doc, '4.1 BUG-4 竞态控制（4 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['C-01', 'epoch 令牌初始为 0', '✅ 通过'],
        ['C-02', 'restartGame 递增 epoch', '✅ 通过'],
        ['C-03', '旧 AI 回调 epoch 不匹配应被丢弃', '✅ 通过'],
        ['C-04', '当前代次回调正常执行', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '4.2 BUG-5 悔棋边界（4 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['C-05', 'moveHistory 为空时不应执行悔棋', '✅ 通过'],
        ['C-06', '玩家走 1 步后（AI 未走）悔 1 步', '✅ 通过'],
        ['C-07', '走 2 步后悔 2 步回到红方', '✅ 通过'],
        ['C-08', '悔棋后棋盘状态恢复', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '4.3 AI 对局与模式流程（6 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['C-09', '完整对局模拟：红走兵 → AI 回手 → 状态恢复', '✅ 通过'],
        ['C-10', 'AI 不会走非法步', '✅ 通过'],
        ['C-11', 'AI 思考超时保护（timeLimit=0 立即返回）', '✅ 通过'],
        ['C-12', 'ban 局 AI 禁子逻辑：随机禁一个非将棋子', '✅ 通过'],
        ['C-13', '心腹局 AI 随机选心腹', '✅ 通过'],
        ['C-14', '内臣（兵）继位后保留 innerKing 额外攻击距离', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

doc.add_page_break()

# ========== 五、E2E 层测试 ==========
add_heading_styled(doc, '五、E2E 浏览器测试详情（17 用例）', level=1)

add_heading_styled(doc, '5.1 页面加载与首屏（4 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['E-01', '页面标题正确', '✅ 通过'],
        ['E-02', 'canvas 存在且尺寸正确（512x568）', '✅ 通过'],
        ['E-03', '首屏棋盘已绘制（非透明像素 > 0）（BUG-1 验证）', '✅ 通过'],
        ['E-04', '状态栏显示红方走棋', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '5.2 经典局对弈（2 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['E-05', '红兵前进 + AI 回手（BUG-2 验证）', '✅ 通过'],
        ['E-06', '悔棋后回到红方回合', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '5.3 模式切换（5 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['E-07', '切到 ban 局：显示请选择要禁用的棋子', '✅ 通过'],
        ['E-08', 'ban 局：红方禁兵后弹窗正常', '✅ 通过'],
        ['E-09', '心腹局：选心腹后弹窗正常', '✅ 通过'],
        ['E-10', '内臣局：选内臣后弹窗正常', '✅ 通过'],
        ['E-11', '切回经典局正常', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '5.4 稳定性与竞态（3 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['E-12', '重新开局正常', '✅ 通过'],
        ['E-13', 'AI 思考期间点重开：无异常且新局正常（BUG-4 验证）', '✅ 通过'],
        ['E-14', '无 JS 控制台错误', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '5.5 ban 进阶局补充（3 用例）', level=2)
make_table(doc,
    ['编号', '测试项', '结果'],
    [
        ['B-01', '切换 banpro：提示选择红方棋子', '✅ 通过'],
        ['B-02', 'banpro：红方禁兵 → AI 禁 → 弹窗 → 开局', '✅ 通过'],
        ['B-03', 'banpro：开局后可正常走棋', '✅ 通过'],
    ]
)
add_para(doc, '', size=Pt(6))

doc.add_page_break()

# ========== 六、修复验证对照 ==========
add_heading_styled(doc, '六、修复验证对照', level=1)
add_para(doc, '对照《修复方案.md》中列出的 9 个 BUG，逐项验证修复效果：')

make_table(doc,
    ['BUG', '问题描述', '级别', '验证结果'],
    [
        ['BUG-1', '首屏棋盘空白，未触发首帧绘制', 'P0', '✅ 已修复（E-03 像素验证通过）'],
        ['BUG-2', 'AI 思考必崩：调用不存在的方法 isCheckmate', 'P0', '✅ 已修复（L-16/E-05 通过）'],
        ['BUG-3', '内臣局被吃方判断反了', 'P1', '✅ 已修复（L-25 通过）'],
        ['BUG-4', 'AI 异步回调与重开/切模式竞态', 'P1', '✅ 已修复（C-01~04/E-13 通过）'],
        ['BUG-5', '悔棋逻辑边界不严谨', 'P1', '✅ 已修复（C-05~08 通过）'],
        ['BUG-6', 'Google Fonts 外网依赖', 'P2', '✅ 已修复（源码确认已替换系统字体）'],
        ['BUG-7', 'favicon 404', 'P2', '✅ 已修复（E-14 无控制台错误）'],
        ['BUG-8', 'ban 局死代码 filter(p=>!game.banBlack)', 'P2', '✅ 已修复（源码确认已修正）'],
        ['BUG-9', '心腹局 render 冗余恒等判断', 'P2', '✅ 已修复（源码确认已简化）'],
    ]
)

add_para(doc, '', size=Pt(6))
add_para(doc, '结论：9 个 BUG 全部修复并通过回归验证，无回退。', size=Pt(12), bold=True, color=RGBColor(0x00, 0x80, 0x00))

doc.add_page_break()

# ========== 七、测试结论 ==========
add_heading_styled(doc, '七、测试结论', level=1)

make_table(doc,
    ['指标', '数值'],
    [
        ['测试用例总数', '60 个'],
        ['通过', '60 个（100%）'],
        ['失败', '0 个'],
        ['覆盖模块', '5 种模式（普通/ban/ban进阶/心腹/内臣）+ AI + 悔棋 + 设置'],
        ['覆盖维度', '规则正确性、AI 稳定性、模式逻辑、边界处理、竞态控制、真实交互、渲染'],
        ['BUG 回归', '9/9 全部通过'],
    ]
)

add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '7.1 质量评估', level=2)
make_table(doc,
    ['维度', '评级', '说明'],
    [
        ['功能完整性', '🟢 优秀', '5 种模式全部可用，走棋规则正确'],
        ['AI 稳定性', '🟢 优秀', '5 档难度搜索无崩溃，走法合法'],
        ['边界处理', '🟢 优秀', '悔棋空历史、竞态、超时均有保护'],
        ['渲染表现', '🟢 优秀', '首屏立即绘制，无空白等待'],
        ['代码质量', '🟡 良好', '9 个历史 BUG 已修复，无新缺陷'],
    ]
)

add_para(doc, '', size=Pt(6))

add_heading_styled(doc, '7.2 建议', level=2)
make_table(doc,
    ['序号', '建议', '优先级'],
    [
        ['1', '保留测试脚本（test_logic.js / test_controller.js / test_e2e.js）供后续回归使用', 'P1'],
        ['2', '关注内臣继位后走法设计（方案 B 已确认，建议补充用户手册说明）', 'P2'],
        ['3', '建议补充单元测试到 CI 流程，防止回归', 'P2'],
    ]
)

add_para(doc, '', size=Pt(12))
add_para(doc, '报告完成', size=Pt(14), bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, 'Luis（露易丝）| 2026-09-04', size=Pt(12), alignment=WD_ALIGN_PARAGRAPH.CENTER)

output_path = r'D:\guodh\file\Loomy\workspace\中国象棋-全量测试报告.docx'
doc.save(output_path)
print('✅ 测试报告已生成: ' + output_path)